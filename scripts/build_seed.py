"""Build the source-registry seed from the team spreadsheet.

Reads `Hack-a-thon 102526.xlsx`, normalizes every agency into the `sources`
schema (supabase/migrations/*_source_registry.sql), adds the shared bid
platforms, and writes:

  data/sources.seed.json   canonical seed, checked into git
  supabase/seed.sql        idempotent upsert, paste into the SQL editor

Usage: python3 scripts/build_seed.py
"""

import json
import re
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "Hack-a-thon 102526.xlsx"
OUT_JSON = ROOT / "data" / "sources.seed.json"
OUT_SQL = ROOT / "supabase" / "seed.sql"
AGENCIES = ROOT / "data" / "agencies"  # OR/ID lists and WA fixes from discovery

# Shared portals: one search covers many agencies. Unverified URLs stay 'candidate'.
PLATFORMS = [
    ("us-federal-sam-gov", "SAM.gov contract opportunities", "US", "federal", "https://sam.gov/", "https://sam.gov/search/?index=opp", "sam", "api"),
    ("wa-platform-webs", "Washington's Electronic Business Solution (WEBS)", "WA", "platform", "https://des.wa.gov/", "https://pr-webs-vendor.des.wa.gov/", "webs", "browser"),
    ("wa-platform-bxwa", "Builders Exchange of Washington", "WA", "platform", "https://www.bxwa.com/", "https://www.bxwa.com/", "bxwa", "browser"),
    ("or-platform-oregonbuys", "OregonBuys", "OR", "platform", "https://oregonbuys.gov/", "https://oregonbuys.gov/bso/", "oregonbuys", "browser"),
    ("id-platform-luma", "Idaho Division of Purchasing (Luma)", "ID", "platform", "https://purchasing.idaho.gov/", "https://purchasing.idaho.gov/", "luma", "browser"),
]

# Known spreadsheet errors: id -> fields to override.
FIXES = {
    "wa-county-adams": {
        "homepage_url": None,
        "procurement_url": None,
        "notes": "Spreadsheet URLs pointed to Adams County, Colorado (adcogov.org); removed.",
    },
    "wa-county-spokane": {
        "notes": "Procurement URL is the City of Spokane portal, not the county's.",
    },
}

# Cities in the accuracy-test counties, used to tag WA rows that have a city but no county.
TEST_COUNTY_CITIES = {
    "King": "Seattle Bellevue Kent Renton Federal-Way Kirkland Auburn Redmond Sammamish Burien Shoreline Issaquah Des-Moines SeaTac Kenmore Tukwila Maple-Valley Covington Mercer-Island Lake-Forest-Park Newcastle Normandy-Park Snoqualmie North-Bend Enumclaw Black-Diamond Duvall Carnation Woodinville Pacific Algona Medina Clyde-Hill Yarrow-Point Hunts-Point Beaux-Arts Skykomish Vashon Bothell",
    "Snohomish": "Everett Marysville Edmonds Lynnwood Lake-Stevens Mukilteo Mill-Creek Mountlake-Terrace Monroe Arlington Snohomish Stanwood Granite-Falls Sultan Gold-Bar Brier Woodway Darrington Index",
    "Pierce": "Tacoma Lakewood Puyallup University-Place Bonney-Lake Gig-Harbor Fife Sumner DuPont Edgewood Fircrest Orting Steilacoom Buckley Eatonville Milton Roy Ruston Wilkeson Carbonado South-Prairie Spanaway Graham",
}
COUNTY_OF_CITY = {c.replace("-", " ").lower(): county for county, cities in TEST_COUNTY_CITIES.items() for c in cities.split()}

PLATFORM_PATTERNS = [
    ("bidnetdirect.com", "bidnet"),
    ("bonfirehub.com", "bonfire"),
    ("procurement.opengov.com", "opengov"),
    ("questcdn.com", "questcdn"),
    ("demandstar.com", "demandstar"),
    ("publicpurchase.com", "publicpurchase"),
    ("ebidexchange", "ebidexchange"),
    ("/bids.aspx", "civicplus"),
    ("/bids/", "custom"),
]


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def clean(s):
    if s is None:
        return None
    s = re.sub(r"\s+", " ", str(s).replace("\xa0", " ").replace("•", "")).strip()
    return s or None


def url_of(cell):
    """A cell's URL, from its hyperlink or its text. Notes like 'call for notes' are dropped."""
    for v in (cell.hyperlink.target if cell.hyperlink else None, cell.value):
        v = clean(v)
        if v and re.match(r"https?://", v):
            return v
    return None


def platform_of(url):
    if not url:
        return None
    low = url.lower()
    for needle, name in PLATFORM_PATTERNS:
        if needle in low:
            return name
    return "custom"


def fetch_method_of(url):
    if not url:
        return None
    return "pdf" if urlparse(url).path.lower().endswith(".pdf") else "html"


def special_type(name):
    n = name.lower()
    for words, t in [
        (("school district", "educational service district", "esd "), "school"),
        (("university", "college"), "higher_ed"),
        (("pud", "public utility", "utility district", "utilities", "energy northwest"), "utility"),
        (("port of", "port authority"), "port"),
        (("hospital", "phd", "health"), "hospital"),
        (("fire",), "fire"),
        (("911", "radio network", "rivercom"), "emergency_comms"),
        (("water", "sewer", "irrigation"), "water"),
        (("transportation", "transit"), "transit"),
        (("park",), "parks"),
    ]:
        if any(w in n for w in words):
            return t
    return "special_district"


def source(id_, name, agency_type, *, city=None, county=None, homepage=None, procurement=None, raw=None):
    return {
        "id": id_,
        "name": name,
        "state": "WA",
        "agency_type": agency_type,
        "county": county,
        "city": city,
        "lat": None,
        "lng": None,
        "homepage_url": homepage,
        "procurement_url": procurement,
        "platform": platform_of(procurement),
        "fetch_method": fetch_method_of(procurement),
        "status": "candidate" if procurement else "missing",
        "last_verified": None,
        "discovered_via": "spreadsheet",
        "notes": None,
        "raw": raw,
    }


def header_index(ws):
    return {clean(c.value): i for i, c in enumerate(ws[1]) if clean(c.value)}


def raw_row(ws, row):
    h = {i: k for k, i in header_index(ws).items()}
    out = {"sheet": ws.title}
    for i, c in enumerate(row):
        v = url_of(c) or clean(c.value)
        if v is not None:
            out[h.get(i, f"col{i}")] = v
    return out


def read_school_districts(ws):
    h = header_index(ws)
    for row in ws.iter_rows(min_row=2):
        name = clean(row[h["DistrictName"]].value)
        if not name:
            continue
        # The 'WEB' column holds the ESD page; the district's own site sits under 'PLAN'.
        yield source(
            f"wa-school-{slug(name)}", name, "school",
            city=(clean(row[h["City"]].value) or "").title() or None,
            homepage=url_of(row[h["PLAN"]]),
            procurement=url_of(row[h["PROCUREMENT"]]),
            raw=raw_row(ws, row),
        )


def read_puds(ws):
    h = header_index(ws)
    for row in ws.iter_rows(min_row=2):
        name = clean(row[1].value)
        if not name or "PUD" not in name.upper() and "UTILIT" not in name.upper():
            continue
        yield source(
            f"wa-utility-{slug(name)}", name.title().replace("Pud", "PUD"), "utility",
            procurement=url_of(row[h["PROCUREMENT"]]),
            raw=raw_row(ws, row),
        )


def read_governments(ws):
    h = header_index(ws)
    for row in ws.iter_rows(min_row=2):
        code, name = row[h["CODE"]].value, clean(row[h["NAME OF AGENCY"]].value)
        if not name:
            continue
        homepage, procurement = url_of(row[h["WEBSITE"]]), url_of(row[h["PROCUREMENT"]])
        raw = raw_row(ws, row)
        if code == 1:
            yield source("wa-state-washington", "State of Washington", "state", homepage=homepage, procurement=procurement, raw=raw)
        elif code == 2:
            county = name.title().replace("Snohomsih", "Snohomish").replace("Kittitias", "Kittitas")
            yield source(f"wa-county-{slug(county)}", f"{county} County", "county", county=county, homepage=homepage, procurement=procurement, raw=raw)
        elif code == 3:
            city = name.title()
            yield source(f"wa-city-{slug(city)}", f"City of {city}", "city", city=city, homepage=homepage, procurement=procurement, raw=raw)
        else:
            t = special_type(name)
            yield source(f"wa-{t.replace('_', '-')}-{slug(name)}", name, t, homepage=homepage, procurement=procurement, raw=raw)


def read_hospitals(ws):
    h = header_index(ws)
    for row in ws.iter_rows(min_row=2):
        name, kind = clean(row[1].value), clean(row[3].value)
        if not name or not kind:  # continuation rows hold phone/fax only
            continue
        m = re.match(r"(.+?),\s*WA", clean(row[5].value) or "")
        yield source(
            f"wa-hospital-{slug(name)}", name, "hospital",
            city=m.group(1) if m else None,
            county=(re.match(r"(.+?) County", name) or [None, None])[1],
            homepage=url_of(row[2]),
            procurement=url_of(row[h["PROCUREMENT"]]),
            raw=raw_row(ws, row),
        )


def domain(url):
    return urlparse(url).netloc.lower().removeprefix("www.") if url else None


def merge(rows):
    """Drop duplicates across sheets, keyed by id and then by homepage domain."""
    by_id, by_domain = {}, {}
    for r in rows:
        d = domain(r["homepage_url"])
        existing = by_id.get(r["id"]) or (by_domain.get(d) if d and r["agency_type"] != "county" else None)
        if existing is None:
            by_id[r["id"]] = r
            if d:
                by_domain.setdefault(d, r)
            continue
        for k, v in r.items():
            if existing.get(k) is None and v is not None:
                existing[k] = v
        if existing["procurement_url"]:
            existing["status"] = "candidate"
    return list(by_id.values())


def sql_literal(v):
    if v is None:
        return "null"
    if isinstance(v, (dict, list)):
        return "'" + json.dumps(v, ensure_ascii=False).replace("'", "''") + "'::jsonb"
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def write_sql(rows):
    cols = list(rows[0].keys())
    updates = ", ".join(f"{c} = coalesce(excluded.{c}, public.sources.{c})" for c in cols if c != "id")
    values = ",\n".join("  (" + ", ".join(sql_literal(r[c]) for c in cols) + ")" for r in rows)
    OUT_SQL.write_text(
        "-- Generated by scripts/build_seed.py. Do not edit by hand.\n"
        f"insert into public.sources ({', '.join(cols)}) values\n{values}\n"
        f"on conflict (id) do update set {updates};\n"
    )


def merge_agency_files(rows):
    """Fold in data/agencies/: `*.json` lists add rows; wa_procurement_fixes.json patches by id."""
    by_id = {r["id"]: r for r in rows}
    template = source("", "", "special_district")
    files = [(p.name, json.loads(p.read_text())) for p in sorted(AGENCIES.glob("*.json"))]
    # New rows from every file first, so an update in one file can patch a row added by another.
    for name, data in files:
        for r in data.get("new", []) if isinstance(data, dict) else data:
            if r["id"] in by_id:
                by_id[r["id"]].update({k: v for k, v in r.items() if v is not None})
            else:
                by_id[r["id"]] = {**template, **r}
    for name, data in files:
        for u in data.get("updates", []) if isinstance(data, dict) else []:
            if u["id"] in by_id:
                by_id[u["id"]].update({k: v for k, v in u.items() if k != "id"})
            else:
                print(f"  {name}: no source {u['id']}, skipped")
    return list(by_id.values())


def main():
    wb = openpyxl.load_workbook(XLSX)
    rows = [
        *read_governments(wb["Sheet1 (2)"]),
        *read_puds(wb["PUDs"]),
        *read_school_districts(wb["WA_School District"]),
        *read_hospitals(wb["Hospitals"]),
    ]
    rows = merge(rows)
    for id_, fix in FIXES.items():
        for r in rows:
            if r["id"] == id_:
                r.update(fix)
                r["platform"] = platform_of(r["procurement_url"])
                r["fetch_method"] = fetch_method_of(r["procurement_url"])
                r["status"] = "candidate" if r["procurement_url"] else "missing"
    for id_, name, state, t, home, proc, plat, method in PLATFORMS:
        r = source(id_, name, t, homepage=home, procurement=proc)
        r.update(state=state, platform=plat, fetch_method=method, discovered_via="manual")
        rows.append(r)
    rows = merge_agency_files(rows)
    for r in rows:
        if r["state"] == "WA" and not r["county"] and r["city"]:
            r["county"] = COUNTY_OF_CITY.get(r["city"].lower())
    rows.sort(key=lambda r: r["id"])

    OUT_JSON.write_text(json.dumps(rows, indent=2, ensure_ascii=False) + "\n")
    write_sql(rows)

    have = sum(1 for r in rows if r["procurement_url"])
    print(f"{len(rows)} sources, {have} with a procurement URL")
    for t in sorted({r["agency_type"] for r in rows}):
        sub = [r for r in rows if r["agency_type"] == t]
        print(f"  {t:17} {len(sub):4}  procurement {sum(1 for r in sub if r['procurement_url'])}")


if __name__ == "__main__":
    main()
